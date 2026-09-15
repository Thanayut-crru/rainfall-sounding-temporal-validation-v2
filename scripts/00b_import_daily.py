"""Convert the horizontal TMD daily-rainfall worksheet to dated records."""
from pathlib import Path
import argparse, calendar, json
import numpy as np
import pandas as pd
from openpyxl import load_workbook

STATIONS={'551201':'surat_thani','551401':'phra_saeng','552201':'nakhon_si_thammarat',
          '552401':'chawang','568401':'sadao','568501':'songkhla','568502':'hat_yai'}

def amount(x):
    if x is None or (isinstance(x,str) and x.strip()=='-'): return np.nan
    if isinstance(x,str) and x.strip().upper()=='T': return 0.0
    x=float(x)
    if x<0: raise ValueError(f'negative rainfall: {x}')
    return x

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--input',type=Path,required=True); ap.add_argument('--output',type=Path,default=Path(__file__).resolve().parent.parent/'data/raw/rainfall_tmd_daily_east.csv'); o=ap.parse_args()
    ws=load_workbook(o.input,read_only=True,data_only=True)['ปริมาณฝนรายวัน']; records=[]
    for row in ws.iter_rows(min_row=6,values_only=True):
        station_text,month=row[1],row[2]
        if not station_text or not hasattr(month,'year'): continue
        code=str(station_text).split('-',1)[0].strip()
        if code not in STATIONS: continue
        ndays=calendar.monthrange(month.year,month.month)[1]
        for day in range(1,ndays+1):
            records.append({'date':pd.Timestamp(month.year,month.month,day),'station':STATIONS[code],
                            'station_code':code,'rain_mm':amount(row[day+2])})
    long=pd.DataFrame(records)
    if long.duplicated(['date','station']).any(): raise ValueError('duplicate station-date in daily worksheet')
    wide=long.pivot(index='date',columns='station',values='rain_mm').add_prefix('RF_').reset_index()
    o.output.parent.mkdir(parents=True,exist_ok=True); wide.to_csv(o.output,index=False)
    qc=long.assign(year=long.date.dt.year,available=long.rain_mm.notna(),rain=lambda x:x.rain_mm>=.1).groupby(['station','station_code','year']).agg(n_days=('date','size'),n_available=('available','sum'),n_rain_days=('rain','sum')).reset_index()
    base=Path(__file__).resolve().parent.parent
    qc.to_csv(base/'outputs/rainfall_daily_east_qc.csv',index=False)
    meta={'worksheet':'ปริมาณฝนรายวัน','layout':'one station-month per row; days 1-31 across columns D-AH','trace_mm':0.0,'stations':STATIONS,'date_min':str(long.date.min().date()),'date_max':str(long.date.max().date()),'n_records':len(long)}
    (base/'outputs/rainfall_daily_east_metadata.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    print(wide[['date','RF_nakhon_si_thammarat']].agg({'date':['min','max'],'RF_nakhon_si_thammarat':'count'})); print(qc[qc.station=='nakhon_si_thammarat'].to_string(index=False))
if __name__=='__main__': main()
